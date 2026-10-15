"""기획 나이아 데이터 테이블을 UE DataTable 임포트용 CSV로 변환한다.

기획 원본은 CSV 또는 xlsx로 받을 수 있다. 두 형식 모두 4행의 컬럼명을 기준으로 읽고,
6행부터 Naia_ID가 있는 행만 변환한다. 기획 컬럼명과 FNUNaiaDefinition 프로퍼티명,
한글 enum·GameplayTag 값의 차이는 이 스크립트가 매핑한다.

사용법:
    python Scripts/ConvertNaiaTable.py "최신 기획 종합본.xlsx" --check

출력: Scripts/Data/NaiaTable.csv

기존 출력 CSV의 그림·모션 경로와 MaxSpawnLevel은 원본에 칸이 없으면 보존한다.
MaxLevel은 성장 상한이므로 랜덤 스폰 상한으로 변환하지 않는다.
원본에서 빠진 기존 행도 보존한다. 의도한 삭제는 --drop-missing으로 지정한다.
--check는 변경 내용을 보여주기만 하고 파일을 쓰지 않는다.
"""

import argparse
import csv
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path


# 기획 문서 구조: 1~3행은 그룹 헤더, 4행은 컬럼명, 5행은 타입 표기, 6행부터 데이터
DEFAULT_SHEET = "나이아 테이블"

HEADER_ROW = 4
FIRST_DATA_ROW = 6

# 기획 컬럼 -> FNUNaiaDefinition 프로퍼티
# 속성만 입력 컬럼 하나를 출력 프로퍼티 두 개로 분리한다.
# 기획 수치는 변환하고, 아래 선택 컬럼의 클라이언트 값은 기존 CSV에서 보존한다.
COLUMN_MAP = [
    ("Naia_ID", "DataIndex", "int"),
    ("Name_Key", "DisplayNameKey", "name"),
    ("Name_KR", "KoreanName", "text"),
    ("Name_EN", "EnglishName", "text"),
    ("분류", "Category", "category"),
    ("속성", ("ElementTag", "SecondaryElementTag"), "elements"),
    ("등급", "Grade", "grade"),
    ("진화 단계", "EvolutionStage", "stage"),
    ("진화 충족 레벨", "EvolutionLevel", "int"),
    ("다음 진화체 ID", "NextEvolutionId", "int"),
    ("포획 가능 여부", "bCapturable", "bool"),
    ("SpawnLevel", "SpawnLevel", "int"),
    ("MaxSpawnLevel", "MaxSpawnLevel", "int"),
    ("BaseHP", "BaseHealth", "float"),
    ("BaseSP", "BaseSP", "float"),
    ("BaseATK", "BaseAttackPower", "float"),
    ("BaseMATK", "BaseMagicAttackPower", "float"),
    ("BaseDEF", "BaseDefense", "float"),
    ("BaseSpeed", "BaseSpeed", "float"),
    ("CritRate", "CriticalRate", "float"),
    ("CritDamage", "CriticalDamage", "float"),
    ("BaseEXP", "BaseRequiredExperience", "int"),
    ("HP_Growth", "HealthGrowth", "float"),
    ("SP_Growth", "SPGrowth", "float"),
    ("ATK_Growth", "AttackGrowth", "float"),
    ("MATK_Growth", "MagicAttackGrowth", "float"),
    ("DEF_Growth", "DefenseGrowth", "float"),
    ("Speed_Growth", "SpeedGrowth", "float"),
    ("EXP_Growth", "RequiredExperienceGrowth", "int"),
    ("EXP", "Experience", "int"),
    ("Gold", "Gold", "int"),
    ("PrimaryDrop1ID", "DropItem1Id", "name"),
    ("PrimaryDrop2ID", "DropItem2Id", "name"),
    ("DropRate", "DropRate", "float"),
    ("Passive1_ID", "PassiveSkill1Id", "int"),
    ("Passive2_ID", "PassiveSkill2Id", "int"),
    ("Passive1_Chance", "PassiveSkill1Chance", "float"),
    ("Passive2_Chance", "PassiveSkill2Chance", "float"),
    ("리소스 경로", "AnimationClass", "class"),
    ("초상화 플립북", "PreviewFlipbook", "class"),
    ("HitAnimation", "HitAnimation", "class"),
    ("AttackAnimation", "AttackAnimation", "class"),
    ("bAttackInPlace", "bAttackInPlace", "bool"),
]

# 클라이언트가 채우는 선택 컬럼. 그림 경로의 빈 칸은 기존 값을 보존하고,
# 명시적인 None / null / - 는 삭제로 취급한다. bool의 FALSE도 명시적 변경이다.
OPTIONAL_SHEET_COLUMNS = {
    "리소스 경로", "초상화 플립북", "HitAnimation", "AttackAnimation",
    "bAttackInPlace", "MaxSpawnLevel",
}
CLIENT_PROPERTIES = {
    "AnimationClass", "PreviewFlipbook", "HitAnimation", "AttackAnimation",
    "bAttackInPlace",
}
COLUMN_ALIASES = {"리소스 경로": "AnimationClass", "초상화 플립북": "PreviewFlipbook"}
# 보스의 미작성 칸을 0으로 바꾸면 정상적인 대련 데이터가 무너진다.
BOSS_REQUIRED_COLUMNS = {
    "BaseHP", "BaseSP", "BaseATK", "BaseMATK", "BaseDEF", "BaseSpeed",
    "CritRate", "CritDamage", "HP_Growth", "SP_Growth", "ATK_Growth",
    "MATK_Growth", "DEF_Growth", "Speed_Growth",
}

CATEGORY = {"스타팅 나이아": "Starter", "야생 정령": "Wild", "보스 나이아": "Boss"}
GRADE = {"일반": "Normal", "희귀": "Rare", "전설": "Legendary"}
STAGE = {"1단계": "Stage1", "2단계": "Stage2", "3단계": "Stage3"}
ELEMENT = {
    "불": "NU.Element.Fire",
    "물": "NU.Element.Water",
    "풀": "NU.Element.Grass",
    "번개": "NU.Element.Lightning",
    "땅": "NU.Element.Earth",
    "바람": "NU.Element.Wind",
}

# 값이 비었다는 뜻으로 기획에서 사용하는 표기
EMPTY_MARKS = {"", "-", "none", "null"}


def is_empty(value):
    return value is None or str(value).strip().lower() in EMPTY_MARKS


class Report:
    """변환 중 사람이 확인해야 할 항목을 모은다."""

    def __init__(self):
        self.unknown_enums = []
        self.dual_elements = []
        self.missing_name_key = []
        self.preserved_cells = 0
        self.preserved_rows = []


def to_int(value):
    if is_empty(value):
        return "0"
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as error:
        raise ValueError(f"정수로 변환할 수 없는 값: {value}") from error
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f"정수 컬럼에 정수가 아닌 값이 있다: {value}")
    return str(int(number))


def to_float(value):
    if is_empty(value):
        return "0"
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as error:
        raise ValueError(f"실수로 변환할 수 없는 값: {value}") from error
    if not number.is_finite():
        raise ValueError(f"유한한 실수가 아닌 값: {value}")
    if number == 0:
        return "0"
    # Decimal 문자열로 유효 숫자를 보존하고 불필요한 소수점 이하 0만 제거한다.
    return format(number.normalize(), "f")


def convert_cell(kind, value, naia_id, prop, report):
    if kind == "int":
        return to_int(value)

    if kind == "float":
        return to_float(value)

    if kind == "bool":
        if is_empty(value):
            return "FALSE"
        key = str(value).strip().lower()
        if key in {"true", "1"}:
            return "TRUE"
        if key in {"false", "0"}:
            return "FALSE"
        raise ValueError(f"bool로 변환할 수 없는 값: {value}")

    if kind in ("name", "text", "class"):
        if is_empty(value):
            if kind == "name" and prop == "DisplayNameKey":
                report.missing_name_key.append(naia_id)
            return ""
        return str(value).strip()

    if kind == "category":
        return lookup(CATEGORY, value, naia_id, prop, report)

    if kind == "grade":
        return lookup(GRADE, value, naia_id, prop, report)

    if kind == "stage":
        return lookup(STAGE, value, naia_id, prop, report)

    if kind == "elements":
        return convert_elements(value, naia_id, report)

    raise ValueError(f"알 수 없는 변환 종류: {kind}")


def lookup(table, value, naia_id, prop, report):
    key = "" if value is None else str(value).strip()
    if key in table:
        return table[key]
    if key in table.values():
        return key
    if is_empty(key) and prop in {"Category", "Grade", "EvolutionStage"}:
        return "None" if prop == "EvolutionStage" else "Unknown"

    report.unknown_enums.append((naia_id, prop, key))
    return ""


def convert_elements(value, naia_id, report):
    """기획 속성을 주 속성과 보조 속성 GameplayTag로 나눈다."""
    if is_empty(value):
        return "", ""

    parts = [part.strip() for part in str(value).split(",") if part.strip()]
    if len(parts) > 1:
        report.dual_elements.append((naia_id, str(value).strip()))
    if len(parts) > 2:
        print(
            f"[경고] {naia_id}: 속성이 세 개 이상이다. 앞의 두 속성만 사용한다: {value}",
            file=sys.stderr,
        )

    primary = lookup(ELEMENT, parts[0], naia_id, "ElementTag", report)
    secondary = ""
    if len(parts) > 1:
        secondary = lookup(
            ELEMENT, parts[1], naia_id, "SecondaryElementTag", report
        )
    return primary, secondary


def read_csv_rows(source_path):
    with source_path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def read_xlsx_rows(source_path, sheet_name=None):
    """워크북에서 시트를 이름으로 골라 읽는다.

    기획이 표를 종합본 하나로 합쳐서, 첫 시트를 읽으면 다른 표를 읽게 된다.
    이름을 못 찾으면 실패한다. 다른 시트를 나이아 원본으로 오인하지 않는다.
    """
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("xlsx 입력에는 openpyxl이 필요하다.") from error

    workbook = openpyxl.load_workbook(source_path, data_only=True, read_only=True)

    try:
        worksheet = next(
            (candidate for candidate in workbook.worksheets
             if not sheet_name or candidate.title.strip() == sheet_name.strip()),
            None,
        )
        if worksheet is None:
            raise ValueError(
                f"'{sheet_name}' 시트가 없다. 이 파일의 시트: {', '.join(workbook.sheetnames)}"
            )
        print(f"[확인] 시트 '{worksheet.title}'")
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_source_rows(source_path, sheet_name=None):
    suffix = source_path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(source_path)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx_rows(source_path, sheet_name)
    raise ValueError(f"지원하지 않는 입력 형식: {source_path.suffix}")


def make_header(rows):
    if len(rows) < HEADER_ROW:
        raise ValueError(f"입력 문서에 {HEADER_ROW}행 헤더가 없다.")

    header = {}
    for column_index, value in enumerate(rows[HEADER_ROW - 1]):
        if is_empty(value):
            continue
        name = str(value).strip()
        if name in header:
            raise ValueError(
                f"중복된 컬럼명: {name}"
                f" ({HEADER_ROW}행 {header[name] + 1}번째 칸과 {column_index + 1}번째 칸)"
                " — 어느 쪽을 쓸지 알 수 없으므로 시트에서 하나를 지워야 한다"
            )
        header[name] = column_index
    return header


def source_value(row, header, column):
    index = header.get(column)
    return row[index] if index is not None and index < len(row) else None


def read_existing(csv_path):
    """기존 클라이언트 값은 ID로 대조한다. 행 순서에 의존하지 않는다."""
    if not csv_path.is_file():
        return [], {}
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = reader.fieldnames or []
        if "Name" not in fields or "DataIndex" not in fields:
            raise ValueError(f"보존 기준 CSV에 Name/DataIndex가 없다: {csv_path}")
        records = {}
        for row in reader:
            if not row.get("Name"):
                continue
            naia_id = int(to_int(row["DataIndex"]))
            if naia_id <= 0 or row["Name"] != str(naia_id) or naia_id in records:
                raise ValueError(f"보존 기준 CSV의 ID/행 이름이 잘못되었다: {row['Name']}")
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"보존 기준 CSV의 컬럼 수가 맞지 않는다: {row['Name']}")
            records[naia_id] = row
        return fields, records


def validate_record(record, source_row, header):
    naia_id = record["Name"]
    if record["Category"] == "Boss":
        missing = sorted(
            column for column in BOSS_REQUIRED_COLUMNS
            if is_empty(source_value(source_row, header, column))
        )
        if missing:
            raise ValueError(f"보스 스탯이 미작성이다(0도 명시해야 함): {', '.join(missing)}")
    spawn = int(record["SpawnLevel"])
    maximum = int(record["MaxSpawnLevel"])
    if spawn < 1 or (maximum != 0 and maximum < spawn):
        raise ValueError(f"스폰 레벨 범위가 잘못되었다: {spawn} ~ {maximum}")
    if naia_id == "1701" and (spawn != 10 or maximum not in (0, 10)):
        raise ValueError("1701 브램블코어 대련은 Lv.10 고정이어야 한다")


def print_changes(existing, rows):
    changes = 0
    for row in rows:
        previous = existing.get(int(row["DataIndex"]))
        if previous is None:
            print(f"       + {row['Name']}: {row.get('KoreanName', '')}")
            changes += 1
            continue
        for field, value in row.items():
            if previous.get(field, "") != value:
                print(f"       {row['Name']}.{field}: {previous.get(field, '')!r} -> {value!r}")
                changes += 1
    removed = set(existing) - {int(row["DataIndex"]) for row in rows}
    for naia_id in sorted(removed):
        print(f"       - {naia_id}")
        changes += 1
    print(f"[확인] 변경 {changes}건 (파일 저장 없음)")


def output_properties():
    properties = []
    for _, prop, _ in COLUMN_MAP:
        if isinstance(prop, tuple):
            properties.extend(prop)
        else:
            properties.append(prop)
    return properties


def convert(source_path, csv_path, sheet_name=DEFAULT_SHEET, *, preserve_from=None,
            check=False, drop_missing=False):
    try:
        source_rows = read_source_rows(source_path, sheet_name)
        header = make_header(source_rows)
        for column, alias in COLUMN_ALIASES.items():
            if column not in header and alias in header:
                header[column] = header[alias]
        old_fields, existing = read_existing(preserve_from or csv_path)
    except (OSError, ValueError, RuntimeError) as error:
        sys.stdout.flush()
        print(f"[실패] 입력 파일을 읽을 수 없다: {error}", file=sys.stderr)
        return 1

    missing = sorted(
        {
            sheet_col
            for sheet_col, _, _ in COLUMN_MAP
            if sheet_col not in header and sheet_col not in OPTIONAL_SHEET_COLUMNS
        }
    )
    if missing:
        print(f"[실패] 기획 문서에 없는 컬럼: {', '.join(missing)}", file=sys.stderr)
        print(
            "       기획 컬럼명이 바뀌었다면 COLUMN_MAP을 갱신해야 한다.",
            file=sys.stderr,
        )
        return 1

    report = Report()
    output_rows = []
    seen_ids = set()
    if "MaxLevel" in header:
        print("[확인] MaxLevel은 성장 상한이다. MaxSpawnLevel로 변환하지 않는다.")

    for row_index, source_row in enumerate(
        source_rows[FIRST_DATA_ROW - 1 :], start=FIRST_DATA_ROW
    ):
        id_index = header["Naia_ID"]
        raw_id = source_row[id_index] if id_index < len(source_row) else None
        if is_empty(raw_id):
            continue

        try:
            naia_id = int(to_int(raw_id))
            if naia_id <= 0:
                raise ValueError("ID는 양수여야 한다")
        except (TypeError, ValueError):
            print(
                f"[실패] Naia_ID가 양의 정수가 아니다: '{raw_id}' ({row_index}행)",
                file=sys.stderr,
            )
            return 1

        if naia_id in seen_ids:
            print(
                f"[실패] Naia_ID 중복: {naia_id} ({row_index}행)",
                file=sys.stderr,
            )
            return 1
        seen_ids.add(naia_id)

        # UE DataTable의 RowName은 Naia_ID 문자열을 사용한다.
        record = dict(existing.get(naia_id, {}))
        record["Name"] = str(naia_id)

        try:
            if naia_id == 9001:
                raise ValueError("대련 보스 9001은 폐기되었다. 브램블코어는 1701을 사용한다")
            if naia_id == 1702 and (
                source_value(source_row, header, "Name_KR") == "할로우코어"
                or str(source_value(source_row, header, "Name_EN")).strip().lower() == "hollowcore"
            ):
                raise ValueError("구버전 할로우코어 1702 행이다. 원본에서 1703으로 옮겨야 한다")
            for sheet_col, prop, kind in COLUMN_MAP:
                value = source_value(source_row, header, sheet_col)
                # 원본에 없거나 아직 작성하지 않은 클라이언트 값은 지우지 않는다.
                if prop in CLIENT_PROPERTIES and (value is None or str(value).strip() == ""):
                    if prop in record:
                        report.preserved_cells += 1
                        continue
                if prop == "MaxSpawnLevel" and sheet_col not in header and prop in record:
                    report.preserved_cells += 1
                    continue
                converted = convert_cell(kind, value, naia_id, prop, report)
                if isinstance(prop, tuple):
                    record.update(zip(prop, converted))
                else:
                    record[prop] = converted
            validate_record(record, source_row, header)
        except (TypeError, ValueError) as error:
            print(
                f"[실패] {naia_id} 변환 중 오류 ({row_index}행): {error}",
                file=sys.stderr,
            )
            return 1

        output_rows.append(record)

    if report.unknown_enums:
        print("[실패] 매핑에 없는 enum/속성이 있다. CSV를 저장하지 않는다.", file=sys.stderr)
        for naia_id, prop, value in report.unknown_enums:
            print(f"       {naia_id}: {prop} = '{value}'", file=sys.stderr)
        return 1
    if not output_rows:
        print("[실패] 원본에 나이아 행이 없다. CSV를 저장하지 않는다.", file=sys.stderr)
        return 1
    if not drop_missing:
        for naia_id, record in existing.items():
            if naia_id not in seen_ids:
                output_rows.append(record)
                report.preserved_rows.append(naia_id)

    # 기존 컬럼 순서와 클라이언트 전용 컬럼도 유지한다.
    field_names = list(old_fields)
    for field in ["Name"] + output_properties():
        if field not in field_names:
            field_names.append(field)
    if check:
        print_report(output_rows, csv_path, report, check=True)
        print_changes(existing, output_rows)
        return 0
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    # UE CSV 임포터용 파일은 UTF-8 BOM과 CRLF로 기록한다.
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=field_names, lineterminator="\r\n")
        writer.writeheader()
        writer.writerows(output_rows)

    print_report(output_rows, csv_path, report)
    return 0


def print_report(rows, csv_path, report, check=False):
    print(f"[{'검사' if check else '완료'}] {len(rows)}행 -> {csv_path}")

    print(f"[보존] 클라이언트 값/스폰 상한 {report.preserved_cells}칸")
    if report.preserved_rows:
        print(f"[보존] 원본에 없는 기존 행: {', '.join(map(str, report.preserved_rows))}")
        print("       의도한 행 삭제라면 --drop-missing으로 다시 변환한다.")

    if report.dual_elements:
        print(
            f"\n[확인] 이중 속성 {len(report.dual_elements)}행 — 두 속성을 별도 컬럼에 보존했다"
        )
        for naia_id, value in report.dual_elements:
            print(f"       {naia_id}: {value}")

    if report.missing_name_key:
        ids = report.missing_name_key
        print(f"\n[확인] Name_Key 미기입 {len(ids)}행 — 표시 이름 키가 비어 있다")
        print(f"       {ids[0]} ~ {ids[-1]}")

    print("\n원본과 보존 CSV에 모두 없으면 구조체 기본값을 사용하는 프로퍼티:")
    print("       Description, AppearanceDescription, FeatureDescription,")
    print("       Height, Weight, IconTexture")


def main():
    # 한국어 Windows 콘솔에서도 리포트 출력이 실패하지 않게 한다.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="기획 나이아 CSV 또는 xlsx 경로")
    parser.add_argument(
        "--out",
        type=Path,
        default=Path(__file__).parent / "Data" / "NaiaTable.csv",
        help="출력 CSV 경로 (기본: Scripts/Data/NaiaTable.csv)",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=f"읽을 워크시트 이름 (기본: {DEFAULT_SHEET}). 종합본에서 시트를 고르는 데 쓴다",
    )
    parser.add_argument(
        "--preserve-from", type=Path,
        help="보존 기준 UE CSV (기본: --out의 기존 파일). 별도 출력 경로를 시험할 때 사용",
    )
    parser.add_argument("--check", action="store_true", help="변경만 출력하고 CSV를 저장하지 않음")
    parser.add_argument("--drop-missing", action="store_true", help="원본에 없는 기존 행을 삭제")
    args = parser.parse_args()

    if not args.source.is_file():
        print(f"[실패] 입력 파일을 찾을 수 없다: {args.source}", file=sys.stderr)
        return 1

    try:
        if args.source.resolve() == args.out.resolve():
            print("[실패] 입력 파일과 출력 파일은 같을 수 없다.", file=sys.stderr)
            return 1
        if args.preserve_from and not args.preserve_from.is_file():
            print(f"[실패] 보존 기준 CSV를 찾을 수 없다: {args.preserve_from}", file=sys.stderr)
            return 1
        if args.preserve_from and args.source.resolve() == args.preserve_from.resolve():
            print("[실패] 보존 기준은 기획 원본이 아니라 UE CSV여야 한다.", file=sys.stderr)
            return 1
    except OSError as error:
        print(f"[실패] 경로를 확인할 수 없다: {error}", file=sys.stderr)
        return 1

    return convert(args.source, args.out, args.sheet, preserve_from=args.preserve_from,
                   check=args.check, drop_missing=args.drop_missing)


if __name__ == "__main__":
    sys.exit(main())
