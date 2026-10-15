"""종합본의 처치 보상 / 나이아 액티브 습득 시트를 공유 CSV로 변환한다.

기획 컬럼명과 단위는 그대로 두고 UE 행 이름용 Name만 앞에 추가한다.
아직 대응하는 C++ RowStruct/실행 로직은 없으므로 이 CSV 생성이 게임 반영은 아니다.
기존 나이아/스킬/아이템 CSV 및 엑셀 원본은 수정하지 않는다.
"""

import argparse
import csv
import io
import math
import os
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

from ConvertNaiaTable import is_empty, to_float, to_int


REWARD_SHEET = "처치 보상 테이블"
LEARN_SHEET = "나이아 액티브 습득 테이블"
REWARD_FIELDS = [
    "RewardPool_ID", "Source_Naia_ID", "RewardType", "PickCount", "Group",
    "GroupWeight", "Ref_ID", "Ref_KR", "Weight", "Rate", "Description",
]
LEARN_FIELDS = [
    "NaiaActive_ID", "Naia_ID", "Active_ID", "Active_KR", "LearnType",
    "MinLevel", "MaxLevel", "Description",
]
REWARD_TYPES = {"RandomItem", "Capture", "FirstClearNaia"}
LEARN_TYPES = {"Signature", "Replace"}
GROUPS = {"제련", "목공", "조제", "직조", "조리"}


def read_records(worksheet, id_column, fields=None):
    """새 시트는 3행 헤더, 기존 정의표는 4행 헤더. ID 컬럼으로 찾는다."""
    rows = list(worksheet.iter_rows(values_only=True))
    candidates = [i for i, row in enumerate(rows[:10]) if id_column in row]
    if len(candidates) != 1:
        raise ValueError(f"{worksheet.title}: {id_column} 헤더를 하나만 찾을 수 있어야 한다")
    header_index = candidates[0]
    header = {}
    for index, value in enumerate(rows[header_index]):
        if is_empty(value):
            continue
        name = str(value).strip()
        if name in header:
            raise ValueError(f"{worksheet.title}: 중복 컬럼 {name}")
        header[name] = index
    if fields is not None and set(header) != set(fields):
        missing = sorted(set(fields) - set(header))
        extra = sorted(set(header) - set(fields))
        raise ValueError(f"{worksheet.title}: 컬럼 불일치 (누락={missing}, 추가={extra})")
    records = []
    seen = set()
    # 헤더 바로 다음은 타입 행이다. 신규 시트의 첫 데이터(5행)를 놓치지 않는다.
    for row_number, row in enumerate(rows[header_index + 2:], header_index + 3):
        record = {name: row[index] if index < len(row) else None
                  for name, index in header.items()}
        if all(is_empty(value) for value in record.values()):
            continue
        try:
            row_id = positive_int(record[id_column], id_column)
        except ValueError as error:
            raise ValueError(f"{worksheet.title} {row_number}행: {error}") from error
        if row_id in seen:
            raise ValueError(f"{worksheet.title} {row_number}행: 중복 ID {row_id}")
        seen.add(row_id)
        record[id_column] = row_id
        records.append(record)
    if not records:
        raise ValueError(f"{worksheet.title}: 데이터 행이 없다")
    return records


def positive_int(value, field):
    if is_empty(value):
        raise ValueError(f"{field} 값이 비었다")
    result = to_int(value)
    if not 1 <= int(result) <= 2147483647:
        raise ValueError(f"{field}는 양의 int32여야 한다: {value}")
    return result


def number(value, field, *, probability=False):
    if is_empty(value):
        raise ValueError(f"{field} 값이 비었다")
    result = to_float(value)
    if not math.isfinite(float(result)) or abs(float(result)) > 3.402823466e38:
        raise ValueError(f"{field} 값이 float 범위를 벗어났다: {value}")
    if (probability and not 0 <= float(result) <= 1) or (
        not probability and float(result) <= 0
    ):
        expected = "0~1 확률" if probability else "양수"
        raise ValueError(f"{field}는 {expected}여야 한다: {value}")
    return result


def text(value):
    return "" if value is None else str(value).strip()


def require_empty(row, fields):
    for field in fields:
        if not is_empty(row[field]):
            raise ValueError(f"{row['RewardType']}에서 사용하지 않는 {field}에 값이 있다")


def validate_rewards(records, naia_ids, item_ids):
    output = []
    group_settings = {}
    pick_counts = {}
    for raw in records:
        row = {field: text(raw[field]) for field in REWARD_FIELDS}
        try:
            for field in ("RewardPool_ID", "Source_Naia_ID", "Ref_ID", "PickCount"):
                row[field] = positive_int(raw[field], field)
            kind = row["RewardType"]
            if kind not in REWARD_TYPES:
                raise ValueError(f"알 수 없는 RewardType: {kind}")
            if row["Source_Naia_ID"] not in naia_ids:
                raise ValueError(f"나이아 표에 없는 Source_Naia_ID: {row['Source_Naia_ID']}")
            references = item_ids if kind == "RandomItem" else naia_ids
            if row["Ref_ID"] not in references:
                raise ValueError(f"참조 표에 없는 Ref_ID: {row['Ref_ID']} ({kind})")
            if kind == "RandomItem":
                if row["Group"] not in GROUPS:
                    raise ValueError(f"알 수 없는 Group: {row['Group']}")
                require_empty(raw, ("Rate",))
                row["Rate"] = ""
                row["GroupWeight"] = number(raw["GroupWeight"], "GroupWeight")
                row["Weight"] = number(raw["Weight"], "Weight")
                key = (row["Source_Naia_ID"], row["Group"])
                setting = row["GroupWeight"]
                if key in group_settings and group_settings[key] != setting:
                    raise ValueError("동일 보스·그룹의 GroupWeight가 서로 다르다")
                group_settings[key] = setting
                source = row["Source_Naia_ID"]
                if source in pick_counts and pick_counts[source] != row["PickCount"]:
                    raise ValueError("동일 보스의 랜덤 풀 PickCount가 서로 다르다")
                pick_counts[source] = row["PickCount"]
            else:
                require_empty(raw, ("Group", "GroupWeight", "Weight"))
                row.update(Group="", GroupWeight="", Weight="")
                row["Rate"] = number(raw["Rate"], "Rate", probability=True)
                if kind == "FirstClearNaia" and row["Rate"] != "1":
                    raise ValueError("FirstClearNaia는 확정 지급이므로 Rate=1이어야 한다")
            output.append(dict(Name=row["RewardPool_ID"], **row))
        except ValueError as error:
            raise ValueError(f"보상 {raw['RewardPool_ID']}: {error}") from error
    return output


def validate_learns(records, naia_ids, skill_ids):
    output = []
    replacements = defaultdict(list)
    seen = set()
    for raw in records:
        row = {field: text(raw[field]) for field in LEARN_FIELDS}
        try:
            for field in ("NaiaActive_ID", "Naia_ID", "Active_ID", "MinLevel", "MaxLevel"):
                row[field] = positive_int(raw[field], field)
            if row["LearnType"] not in LEARN_TYPES:
                raise ValueError(f"알 수 없는 LearnType: {row['LearnType']}")
            if row["Naia_ID"] not in naia_ids:
                raise ValueError(f"나이아 표에 없는 Naia_ID: {row['Naia_ID']}")
            if row["Active_ID"] not in skill_ids:
                raise ValueError(f"액티브 스킬 표에 없는 Active_ID: {row['Active_ID']}")
            low, high = int(row["MinLevel"]), int(row["MaxLevel"])
            if low > high:
                raise ValueError(f"레벨 범위가 뒤집혔다: {low}~{high}")
            key = (row["Naia_ID"], row["LearnType"], row["Active_ID"], low, high)
            if key in seen:
                raise ValueError("같은 습득 규칙이 중복되었다")
            seen.add(key)
            if row["LearnType"] == "Replace":
                replacements[row["Naia_ID"]].append((low, high, row["NaiaActive_ID"]))
            output.append(dict(Name=row["NaiaActive_ID"], **row))
        except ValueError as error:
            raise ValueError(f"습득 {raw['NaiaActive_ID']}: {error}") from error
    # Signature와 Replace는 동시에 보유할 수 있으므로 서로의 구간은 비교하지 않는다.
    for naia_id, ranges in replacements.items():
        ranges.sort()
        for previous, current in zip(ranges, ranges[1:]):
            if current[0] <= previous[1]:
                raise ValueError(f"나이아 {naia_id}: Replace 구간 중복 {previous[2]} / {current[2]}")
    return output


def load_tables(source):
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("xlsx 입력에는 openpyxl이 필요하다") from error
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        required = (REWARD_SHEET, LEARN_SHEET, "나이아 테이블", "액티브 스킬 테이블", "아이템 테이블")
        missing = [name for name in required if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"필요한 시트가 없다: {', '.join(missing)}")
        naia_ids = {r["Naia_ID"] for r in read_records(workbook["나이아 테이블"], "Naia_ID")}
        skill_ids = {r["Active_ID"] for r in read_records(workbook["액티브 스킬 테이블"], "Active_ID")}
        item_ids = {r["Item_ID"] for r in read_records(workbook["아이템 테이블"], "Item_ID")}
        rewards = validate_rewards(
            read_records(workbook[REWARD_SHEET], "RewardPool_ID", REWARD_FIELDS), naia_ids, item_ids)
        learns = validate_learns(
            read_records(workbook[LEARN_SHEET], "NaiaActive_ID", LEARN_FIELDS), naia_ids, skill_ids)
        return rewards, learns
    finally:
        workbook.close()


def csv_bytes(rows, fields):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=["Name"] + fields, lineterminator="\r\n")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def print_summary(rewards, learns):
    for source in sorted({r["Source_Naia_ID"] for r in rewards}, key=int):
        rows = [r for r in rewards if r["Source_Naia_ID"] == source]
        counts = {kind: sum(r["RewardType"] == kind for r in rows) for kind in sorted(REWARD_TYPES)}
        groups = {r["Group"]: r["GroupWeight"] for r in rows if r["RewardType"] == "RandomItem"}
        print(f"[확인] 나이아 {source} 보상 {counts}, 그룹 가중치 {groups}")
    print(f"[확인] 처치 보상 {len(rewards)}행 / 액티브 습득 {len(learns)}행")
    print("[주의] CSV 생성만 수행한다. 대응 RowStruct·보상 지급·스킬 습득 연결은 후속 작업이다.")


def convert(source, reward_out, learn_out, *, check=False):
    source, reward_out, learn_out = map(Path, (source, reward_out, learn_out))
    staged = []
    try:
        paths = [p.resolve() for p in (source, reward_out, learn_out)]
        if len(set(paths)) != 3:
            raise ValueError("입력 / 출력 파일 경로가 서로 겹친다")
        if source.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ValueError("종합 엑셀(.xlsx/.xlsm)을 입력해야 한다")
        rewards, learns = load_tables(source)
        # 두 표 모두 검증한 뒤에만 출력한다. 한 표의 오류가 다른 표를 덮어쓰지 않는다.
        outputs = [(reward_out, csv_bytes(rewards, REWARD_FIELDS)),
                   (learn_out, csv_bytes(learns, LEARN_FIELDS))]
        print_summary(rewards, learns)
        for path, payload in outputs:
            changed = not path.is_file() or path.read_bytes() != payload
            print(f"[{'변경' if changed else '동일'}] {path}")
        if check:
            print("[확인] --check: 파일 저장 없음")
            return 0
        # 파일별 교체는 atomic. 검증/임시 파일 작성 실패 시 기존 CSV는 유지한다.
        for path, payload in outputs:
            path.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(dir=path.parent, prefix=path.name + ".", delete=False) as handle:
                staged.append((Path(handle.name), path))
                handle.write(payload)
        for temporary, path in staged:
            os.replace(temporary, path)
        return 0
    except (ValueError, RuntimeError, OSError) as error:
        print(f"[실패] {error}", file=sys.stderr)
        return 1
    finally:
        for temporary, _ in staged:
            if temporary.exists():
                temporary.unlink()


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    data = Path(__file__).resolve().parent / "Data"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="기획 종합 엑셀 경로")
    parser.add_argument("--reward-out", type=Path, default=data / "NaiaRewardPoolTable.csv")
    parser.add_argument("--learn-out", type=Path, default=data / "NaiaActiveLearnTable.csv")
    parser.add_argument("--check", action="store_true", help="검증 및 차이 유무만 표시, 파일 저장 안 함")
    args = parser.parse_args()
    return convert(args.source, args.reward_out, args.learn_out, check=args.check)


if __name__ == "__main__":
    sys.exit(main())
